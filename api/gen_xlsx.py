#!/usr/bin/env python3
"""
Generador XLSX para UNI-VIRTUAL
Recibe JSON por stdin, devuelve XLSX binario por stdout
Llamado desde PHP via proc_open()
"""
import sys, json, io, base64
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def hex_color(h):
    return h.lstrip('#')

def build_border():
    s = Side(style='thin', color='DDDDDD')
    return Border(left=s, right=s, top=s, bottom=s)

def apply_header(cell, bg, fg='FFFFFF', size=10, bold=True):
    cell.font = Font(bold=bold, size=size, color=fg, name='Arial')
    cell.fill = PatternFill('solid', fgColor=bg)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def gen_notas(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Calificaciones"
    NAVY, BLUE = "0D1F4E", "2462B0"
    WHITE, LIGHT = "FFFFFF", "EEF3FB"
    GREEN, RED = "1A7A48", "C0392B"
    border = build_border()

    # Fila 1 banner
    ws.merge_cells('A1:K1')
    ws['A1'] = 'UNI-VIRTUAL — Institución Universitaria Antonio José Camacho'
    apply_header(ws['A1'], NAVY, size=14)
    ws.row_dimensions[1].height = 30

    # Fila 2 subtítulo
    ws.merge_cells('A2:K2')
    nombre = data.get('materia', '')
    codigo = data.get('codigo', '')
    ws['A2'] = f'Reporte de Calificaciones — {nombre} ({codigo}) — Semestre 2026-I'
    apply_header(ws['A2'], BLUE, size=11, bold=False)
    ws.row_dimensions[2].height = 22

    # Fila 3 info
    ws.merge_cells('A3:K3')
    ws['A3'] = f"Profesor: {data.get('prof','')}   |   Generado: {data.get('fecha','')}"
    ws['A3'].font = Font(size=10, color='666666', name='Arial')
    ws['A3'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[3].height = 16

    # Fila 4 headers
    headers = ['Nombre','Código','Sem.','Parcial 1','Parcial 2','Talleres',
               'Promedio','Nota Final','Asistencia %','Entregas','Estado']
    for i, h in enumerate(headers):
        cell = ws[f'{get_column_letter(i+1)}4']
        cell.value = h
        apply_header(cell, NAVY)
        cell.border = border
    ws.row_dimensions[4].height = 22
    ws.freeze_panes = 'A5'

    # Datos
    proms_grupo = []
    for ri, e in enumerate(data.get('estudiantes', [])):
        row = ri + 5
        bg = WHITE if ri % 2 == 0 else LIGHT
        notas_vals = [v for v in [e.get('p1'), e.get('p2'), e.get('tall')] if v is not None]
        prom = round(sum(notas_vals)/len(notas_vals), 2) if notas_vals else None
        final = e.get('final') if e.get('final') is not None else prom
        if final is not None:
            proms_grupo.append(float(final))
        estado = 'Aprobado' if final and float(final) >= 3.0 else ('Reprobado' if final else 'En curso')

        vals = [
            e.get('nombre',''), e.get('codigo',''), e.get('semestre', ''),
            e.get('p1'), e.get('p2'), e.get('tall'),
            prom, e.get('final'),
            f"{e.get('asistencia',0)}%", e.get('entregas',''),
            estado
        ]
        for ci, val in enumerate(vals):
            cell = ws[f'{get_column_letter(ci+1)}{row}']
            cell.value = val
            cell.border = border
            cell.fill = PatternFill('solid', fgColor=bg)
            cell.alignment = Alignment(horizontal='left' if ci==0 else 'center', vertical='center')
            if ci in [3,4,5,6,7] and val is not None:
                try:
                    color = GREEN if float(val) >= 3.0 else RED
                    cell.font = Font(bold=True, size=10, color=color, name='Arial')
                except: cell.font = Font(size=10, name='Arial')
            elif ci == 10:
                color = GREEN if estado=='Aprobado' else (RED if estado=='Reprobado' else 'B7860B')
                cell.font = Font(bold=True, size=10, color=color, name='Arial')
                cell.fill = PatternFill('solid', fgColor='E8F5EE' if estado=='Aprobado' else ('FDECEA' if estado=='Reprobado' else 'FDF3E4'))
            else:
                cell.font = Font(size=10, name='Arial')
        ws.row_dimensions[row].height = 18

    # Anchos
    for col, w in zip('ABCDEFGHIJK', [28,14,6,10,10,10,10,11,12,10,12]):
        ws.column_dimensions[col].width = w

    # Hoja métricas
    ws2 = wb.create_sheet("Métricas")
    ws2.merge_cells('A1:B1')
    ws2['A1'] = 'UNI-VIRTUAL — Métricas del Grupo'
    apply_header(ws2['A1'], NAVY, size=14)
    ws2.row_dimensions[1].height = 28

    pg = round(sum(proms_grupo)/len(proms_grupo), 2) if proms_grupo else 0
    apr = sum(1 for p in proms_grupo if p >= 3.0)
    rep = len(proms_grupo) - apr
    pct_apr = round(apr/len(proms_grupo)*100) if proms_grupo else 0

    metricas = [
        ('Materia', nombre), ('Código', codigo),
        ('Profesor', data.get('prof','')), ('Semestre', '2026-I'),
        ('Generado', data.get('fecha','')), ('',''),
        ('Total estudiantes', str(len(data.get('estudiantes',[])))),
        ('Aprobados', str(apr)), ('Reprobados', str(rep)),
        ('Promedio del grupo', str(pg)),
        ('Tasa de aprobación', f"{pct_apr}%"),
    ]
    for i, (k, v) in enumerate(metricas):
        r = i + 2
        bg = "EEF3FB" if i % 2 == 0 else "FFFFFF"
        ws2[f'A{r}'] = k
        ws2[f'A{r}'].font = Font(bold=True, size=11, color=NAVY, name='Arial')
        ws2[f'A{r}'].fill = PatternFill('solid', fgColor=bg)
        ws2[f'B{r}'] = v
        ws2[f'B{r}'].font = Font(size=11, color=BLUE, name='Arial')
        ws2[f'B{r}'].fill = PatternFill('solid', fgColor=bg)
        ws2.row_dimensions[r].height = 22
    ws2.column_dimensions['A'].width = 26
    ws2.column_dimensions['B'].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def gen_asistencia(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia"
    NAVY, BLUE = "0D1F4E", "2462B0"
    WHITE, LIGHT = "FFFFFF", "EEF3FB"
    GREEN, RED = "1A7A48", "C0392B"
    border = build_border()
    fechas = data.get('fechas', [])
    estudiantes = data.get('estudiantes', [])
    total_cols = 2 + len(fechas) + 3  # nombre, codigo, fechas, asistidas, %, estado

    last_col = get_column_letter(total_cols)

    # Banner
    ws.merge_cells(f'A1:{last_col}1')
    ws['A1'] = 'UNI-VIRTUAL — Institución Universitaria Antonio José Camacho'
    apply_header(ws['A1'], NAVY, size=14)
    ws.row_dimensions[1].height = 30

    ws.merge_cells(f'A2:{last_col}2')
    materia = data.get('materia','')
    codigo = data.get('codigo','')
    ws['A2'] = f'Reporte de Asistencia — {materia} ({codigo}) — {data.get("desde","")} al {data.get("hasta","")}'
    apply_header(ws['A2'], BLUE, size=11, bold=False)
    ws.row_dimensions[2].height = 22

    ws.merge_cells(f'A3:{last_col}3')
    ws['A3'] = f"Profesor: {data.get('prof','')}   |   Generado: {data.get('fecha','')}"
    ws['A3'].font = Font(size=10, color='666666', name='Arial')
    ws['A3'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[3].height = 16

    # Headers
    hdrs = ['Nombre', 'Código'] + [f['label'] for f in fechas] + ['Asistidas', '%', 'Estado']
    for i, h in enumerate(hdrs):
        cell = ws[f'{get_column_letter(i+1)}4']
        cell.value = h
        apply_header(cell, NAVY)
        cell.border = border
    ws.row_dimensions[4].height = 22
    ws.freeze_panes = 'A5'

    for ri, e in enumerate(estudiantes):
        row = ri + 5
        bg = WHITE if ri % 2 == 0 else LIGHT
        asist_map = e.get('asistencia', {})
        asistidas = 0
        total_f = len(fechas)

        row_vals = [e.get('nombre',''), e.get('codigo','')]
        for f in fechas:
            v = asist_map.get(f['fecha'])
            if v is True: asistidas += 1
            row_vals.append('S' if v is True else ('N' if v is False else '—'))

        pct = round(asistidas/total_f*100) if total_f > 0 else 0
        estado = 'Cumple' if pct >= 75 else 'Riesgo'
        row_vals += [f'{asistidas}/{total_f}', pct, estado]

        for ci, val in enumerate(row_vals):
            cell = ws[f'{get_column_letter(ci+1)}{row}']
            cell.value = val
            cell.border = border
            cell.fill = PatternFill('solid', fgColor=bg)
            cell.alignment = Alignment(horizontal='left' if ci==0 else 'center', vertical='center')
            if ci >= 2 and ci < 2+len(fechas):
                color = GREEN if val == 'S' else (RED if val == 'N' else '999999')
                cell.font = Font(bold=val in ('S','N'), size=10, color=color, name='Arial')
            elif ci == len(row_vals)-1:  # estado
                color = GREEN if estado=='Cumple' else RED
                cell.font = Font(bold=True, size=10, color=color, name='Arial')
                cell.fill = PatternFill('solid', fgColor='E8F5EE' if estado=='Cumple' else 'FDECEA')
            elif ci == len(row_vals)-2:  # pct
                color = GREEN if pct >= 75 else RED
                cell.font = Font(bold=True, size=10, color=color, name='Arial')
            else:
                cell.font = Font(size=10, name='Arial')
        ws.row_dimensions[row].height = 18

    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 14
    for i in range(len(fechas)):
        ws.column_dimensions[get_column_letter(i+3)].width = 8
    ws.column_dimensions[get_column_letter(3+len(fechas))].width = 10
    ws.column_dimensions[get_column_letter(4+len(fechas))].width = 8
    ws.column_dimensions[get_column_letter(5+len(fechas))].width = 10

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# Main
try:
    payload = json.loads(sys.stdin.read())
    tipo = payload.get('tipo', 'notas')
    if tipo == 'notas':
        xlsx_data = gen_notas(payload)
    else:
        xlsx_data = gen_asistencia(payload)
    # Escribir bytes a stdout
    sys.stdout.buffer.write(xlsx_data)
except Exception as e:
    sys.stderr.write(f"ERROR: {e}\n")
    import traceback
    sys.stderr.write(traceback.format_exc())
    sys.exit(1)
